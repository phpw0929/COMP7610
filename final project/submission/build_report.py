"""Build an editable academic report. Export its DOCX to PDF with Word or the verified renderer."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
FONT = 'Times New Roman'
doc = Document()


def style_font(style, size, bold=False, italic=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if 'theme' in key.lower():
            del fonts.attrib[key]
    for item in style.element.xpath('./w:pPr/w:pBdr | ./w:pPr/w:numPr | ./w:rPr/w:spacing'):
        item.getparent().remove(item)
    for kind in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + kind), FONT)


for style_name, size, bold in [('Normal', 12, False), ('Title', 24, True), ('Subtitle', 14, False), ('Heading 1', 14, True), ('Heading 2', 12, True), ('Caption', 10.5, False)]:
    style_font(doc.styles[style_name], size, bold)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.5
normal.space_before = Pt(0)
normal.space_after = Pt(6)
normal.widow_control = True
for name in ('Heading 1', 'Heading 2'):
    p = doc.styles[name].paragraph_format
    p.space_before = Pt(16)
    p.space_after = Pt(8)
    p.keep_with_next = True
    p.keep_together = True
    p.line_spacing = 1.15
caption = doc.styles['Caption'].paragraph_format
caption.line_spacing = 1.15
caption.space_before = Pt(5)
caption.space_after = Pt(10)
style_font(doc.styles['Caption'], 10.5, italic=True)
for name in ('Header', 'Footer'):
    style_font(doc.styles[name], 9)
    doc.styles[name].paragraph_format.line_spacing = 1


def setup_page(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.45)


setup_page(doc.sections[0])
cover = doc.add_paragraph()
cover.paragraph_format.space_before = Pt(56)
cover.paragraph_format.space_after = Pt(22)
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.add_run('The University of Hong Kong').bold = True
p = doc.add_paragraph('COMP7610', style='Subtitle')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph('Introduction to Web3 and Blockchain')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(36)
p = doc.add_paragraph('Volunteer Service Hour\nVerification DApp', style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(20)
p = doc.add_paragraph('Project Report', style='Subtitle')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for text in ('Draft for team review', '28 October 2026', 'Team names and student IDs to be completed'):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
p = doc.add_paragraph('Implementation and local acceptance are complete.\nFinal signed Sepolia UI evidence and team contributions remain pending.')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(28)
for r in p.runs:
    r.font.size = Pt(11)

body = doc.add_section(WD_SECTION_START.NEW_PAGE)
setup_page(body)
body.header.is_linked_to_previous = False
body.footer.is_linked_to_previous = False
hp = body.header.paragraphs[0]
hp.text = 'COMP7610  |  Volunteer Service Hour Verification'
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
page_num = OxmlElement('w:pgNumType')
page_num.set(qn('w:start'), '1')
body._sectPr.append(page_num)
fp = body.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('Page ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run = OxmlElement('w:r')
text = OxmlElement('w:t')
text.text = '1'
run.append(text)
field.append(run)
fp._p.append(field)


def add_architecture():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1
    p.add_run().add_picture(str(ROOT / 'figures/architecture.png'), width=Inches(6.2))
    inline = p.runs[0]._r.xpath('.//wp:docPr')[0]
    inline.set('descr', 'React uses independent read-only RPC for public queries, MetaMask for signed writes, and checks exact metadata byte hashes.')
    cp = doc.add_paragraph('Figure 1  System architecture and separation of public reads and signed writes', style='Caption')
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_results_table():
    cp = doc.add_paragraph('Table 1  Verification results and evidence boundaries', style='Caption')
    cp.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.35), Inches(1.55), Inches(2.36)]
    headers = ['Verification', 'Environment', 'Observed result']
    records = [
        ['Automated tests', 'Fresh local EVM', '15 passed and 0 failed'],
        ['UI workflow', 'Local EVM', '14 successful transactions'],
        ['Wallet free record lookup', 'Sepolia', 'Real records read successfully'],
        ['New signed UI workflow', 'Sepolia', 'Pending final team run'],
    ]
    for j, label in enumerate(headers):
        table.rows[0].cells[j].text = label
    for record in records:
        cells = table.add_row().cells
        for j, value in enumerate(record):
            cells[j].text = value
    header_flag = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(header_flag)
    for j, width in enumerate(widths):
        table.columns[j].width = width
    for i, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            fill = OxmlElement('w:shd')
            fill.set(qn('w:fill'), 'E8E8E8' if i == 0 else ('F7F7F7' if i % 2 == 0 else 'FFFFFF'))
            tc_pr.append(fill)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                edge = OxmlElement('w:' + side)
                edge.set(qn('w:w'), '90')
                edge.set(qn('w:type'), 'dxa')
                margins.append(edge)
            tc_pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                if j > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = FONT
                    r.font.size = Pt(10.5)
                    r.bold = i == 0
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + side)
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '4')
        edge.set(qn('w:color'), 'D9D9D9')
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


blocks = [part.strip() for part in (ROOT / 'project_report_draft.md').read_text().split('\n\n') if part.strip()]
section = None
body_started = False
paragraph_in_section = 0
for block in blocks:
    if block.startswith('# '):
        continue
    if block.startswith('COMP7610 ') or block.startswith('Implementation and evidence draft'):
        continue
    if block.startswith('## '):
        label = re.sub(r'^(\d+)\.\s*', r'\1 ', block[3:])
        doc.add_heading(label, level=1)
        section = block[3:]
        paragraph_in_section = 0
        body_started = True
        continue
    if not body_started:
        doc.add_heading('Executive summary', level=1)
        body_started = True
    if section == 'References':
        for line in block.splitlines():
            p = doc.add_paragraph(line)
            p.paragraph_format.left_indent = Inches(.25)
            p.paragraph_format.first_line_indent = Inches(-.25)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.size = Pt(10.5)
        continue
    p = doc.add_paragraph(block)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph_in_section += 1
    if section == '4. Implementation' and paragraph_in_section == 1:
        add_architecture()
    if section == '5. Testing and observed results' and paragraph_in_section == 1:
        add_results_table()

props = doc.core_properties
props.title = 'Volunteer Service Hour Verification DApp'
props.subject = 'COMP7610 Project Report Draft'
props.author = 'COMP7610 project team'
props.keywords = 'Draft, volunteer service records, DApp'
props.comments = ''
doc.save(ROOT / 'project_report_draft.docx')
print(ROOT / 'project_report_draft.docx')
